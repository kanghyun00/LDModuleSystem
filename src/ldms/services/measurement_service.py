# -*- coding: utf-8 -*-
from __future__ import annotations

import re
from pathlib import Path
from datetime import datetime
from typing import Optional, Dict

import openpyxl

from src.ldms.process_flow import DEFAULT_MODEL_CODE, normalize_model_code
from src.ldms.app_config import load_path_settings, normalize_and_fill_defaults


# ✅ 모듈번호: 알파벳(1~4) + 숫자(1~4)  ex) A1, A002, B1234, E9
MOD_PAT = re.compile(r"([A-Za-z]{1,4}\d{1,4})")
MODEL_PAT = re.compile(r"(F\d{3,4}\d{1,4}W)", re.IGNORECASE)


def _extract_model_code_from_path(path: Path, default: str = DEFAULT_MODEL_CODE) -> str:
    """
    파일명/상위 폴더명에서 F976370W, F976550W 같은 모델명을 추출.
    없으면 기본 F976370W.
    """
    try:
        s = str(path)
        m = MODEL_PAT.search(s)
        if m:
            return normalize_model_code(m.group(1))
    except Exception:
        pass
    return normalize_model_code(default or DEFAULT_MODEL_CODE)




# ✅ 기본 RAW 저장/조회 루트
# Parameter 경로가 비어있을 때만 현재 프로젝트 기준으로 사용
DEFAULT_RAW_ROOT = Path(__file__).resolve().parents[3] / "data" / "measurements" / "raw"


def ensure_default_raw_root() -> Path:
    """
    기본 RAW 루트 폴더 생성 보장
    """
    DEFAULT_RAW_ROOT.mkdir(parents=True, exist_ok=True)
    return DEFAULT_RAW_ROOT


def _to_float(v):
    if v is None:
        return None
    try:
        return float(v)
    except Exception:
        try:
            return float(str(v).strip())
        except Exception:
            return None


def _extract_module_no_from_filename(path: Path) -> str:
    """
    파일명에서 모듈번호(알파벳+숫자)만 뽑고 나머지는 무시.
    예) e92(리페어...).xlsx -> E92
        A1_test.xlsx -> A1
    """
    stem = path.stem.strip()
    m = MOD_PAT.search(stem)
    return m.group(1).upper() if m else stem.upper()


def _norm(s: str) -> str:
    t = str(s).strip().lower()
    t = t.replace("\n", " ").replace("\r", " ")
    t = re.sub(r"\s+", " ", t)
    t = t.replace("（", "(").replace("）", ")")
    t = t.replace("㎚", "nm")
    return t


def _canon_col(name: str) -> Optional[str]:
    """
    다양한 Summary 헤더를 '표준 의미'로 매핑.
    표준 의미:
      Current, Peak1, Peak2, FWHM(nm), SMSR, Power, Voltage,
      Lid(CH5), Base(CH3), Fiber(CH2), 효율
    """
    n = _norm(name)
    nn = n.replace(" ", "")

    # Current
    if n in ("current", "current(a)", "i", "i(a)", "current (a)"):
        return "Current"

    # Peak
    if "peak1" in n or "peak 1" in n:
        return "Peak1"
    if "peak2" in n or "peak 2" in n:
        return "Peak2"

    # FWHM
    if "fwhm" in n or "반치폭" in n:
        return "FWHM(nm)"

    # SMSR
    if "smsr" in n or "dbm diff" in n or "dbm_diff" in n or "dbmdiff" in n:
        return "SMSR"

    # Power / Voltage
    if n.startswith("power") or n in ("p", "power(w)", "power (w)"):
        return "Power"
    if n.startswith("voltage") or n in ("v", "voltage(v)", "voltage (v)"):
        return "Voltage"

    # Temp Channels
    # 최신: Lid(CH5), Base(CH3), Fiber(CH2)
    # 구버전: Base(CH5) == Lid(CH5)
    # 구버전: PKG(CH3) == Base(CH3)
    # "온도" 단독은 CH3로 간주 -> Base(CH3)
    if nn in ("lid(ch5)", "lidch5", "ch5lid"):
        return "Lid(CH5)"
    if nn in ("base(ch5)", "basech5", "ch5"):
        return "Lid(CH5)"  # 구버전 base(ch5) -> 최신 lid(ch5)

    if nn in ("base(ch3)", "basech3", "ch3"):
        return "Base(CH3)"
    if nn in ("pkg(ch3)", "pkgch3", "pkg"):
        return "Base(CH3)"  # 구버전 pkg(ch3) -> 최신 base(ch3)

    if nn in ("온도", "온도(°c)", "온도(℃)", "temperature", "temp", "temp(°c)", "temp(℃)"):
        return "Base(CH3)"  # 요청: 온도 단독은 CH3 취급

    if nn in ("fiber(ch2)", "fiberch2", "ch2"):
        return "Fiber(CH2)"

    if nn in ("eff", "efficiency", "효율"):
        return "효율"

    return None


def _summary_header_map(ws) -> Dict[int, str]:
    """
    Summary 시트에서 1행을 읽어 표준 헤더 매핑(col_index -> std_name)
    """
    col_map: Dict[int, str] = {}
    max_col = ws.max_column or 0
    for c in range(1, max_col + 1):
        v = ws.cell(1, c).value
        if v is None:
            continue
        std = _canon_col(v)
        if std:
            col_map[c] = std
    return col_map


def _read_summary_points(ws) -> list:
    """
    Summary 시트에서 유도리있게 읽어서 "표준 의미" dict list 반환.
    """
    col_map = _summary_header_map(ws)
    if not col_map:
        return []

    # Current 컬럼 찾기
    cur_col = None
    for c, std in col_map.items():
        if std == "Current":
            cur_col = c
            break
    if cur_col is None:
        return []

    out = []
    for r in range(2, (ws.max_row or 0) + 1):
        current = _to_float(ws.cell(r, cur_col).value)
        if current is None:
            continue

        row_std = {
            "Current": current,
            "Peak1": None,
            "Peak2": None,
            "FWHM(nm)": None,
            "SMSR": None,
            "Power": None,
            "Voltage": None,
            "Lid(CH5)": None,
            "Base(CH3)": None,
            "Fiber(CH2)": None,
            "효율": None,
        }

        for c, std in col_map.items():
            if std == "Current":
                continue
            row_std[std] = _to_float(ws.cell(r, c).value)

        out.append(row_std)
    return out


def _std_to_db_point(std: dict) -> dict:
    """
    표준 의미(dict)를 DB 저장 컬럼(dict)로 변환.
    ✅ DB 스키마는 유지:
      - Lid(CH5) -> base_ch5
      - Base(CH3) -> pkg_ch3
    """
    return {
        "current": std.get("Current"),
        "peak1": std.get("Peak1"),
        "peak2": std.get("Peak2"),
        "fwhm_nm": std.get("FWHM(nm)"),
        "smsr": std.get("SMSR"),
        "power": std.get("Power"),
        "voltage": std.get("Voltage"),
        "base_ch5": std.get("Lid(CH5)"),
        "pkg_ch3": std.get("Base(CH3)"),
        "fiber_ch2": std.get("Fiber(CH2)"),
        "eff": std.get("효율"),
    }


def import_measurement_xlsx(repo, xlsx_path: str, stage: str = "", group_name: str = "", model: str = "") -> int:
    """
    측정 엑셀 1개를 DB로 import
    - stage: stage 폴더명
    - group_name: stage 하위 폴더명(예: 2025-e / 2026-A(1월))
    - model: 모델명. 비워두면 파일명/경로에서 F976370W 같은 문자열을 자동 추출
    return: measurement_files.id
    """
    p = Path(xlsx_path)
    if not p.exists():
        raise FileNotFoundError(str(p))

    module_no = _extract_module_no_from_filename(p)
    file_name = p.name
    file_path = str(p.resolve())
    imported_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    stage = (stage or "").strip()
    group_name = (group_name or "").strip()
    model = normalize_model_code(model or _extract_model_code_from_path(p))

    # ✅ 측정 파일 import 시 모듈번호-모델 매핑도 함께 보정
    try:
        if module_no and hasattr(repo, "upsert_module"):
            repo.upsert_module(module_no=module_no, model=model)
    except Exception:
        pass

    measurement_id = repo.upsert_measurement_file(
        module_no=module_no,
        stage=stage,
        group_name=group_name,
        file_name=file_name,
        file_path=file_path,
        imported_at=imported_at,
    )

    wb = openpyxl.load_workbook(file_path, data_only=True)

    # 재import 대비: 기존 포인트 삭제
    repo.delete_measurement_points(measurement_id)

    # Summary (버전 유도리)
    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        std_points = _read_summary_points(ws)
        for std in std_points:
            repo.insert_measurement_summary_point(measurement_id, _std_to_db_point(std))

    # Spectrum (16A / 20A)
    for sheet in ("16A Chart", "20A Chart"):
        if sheet not in wb.sheetnames:
            continue

        ws = wb[sheet]
        rows = []
        for r in range(2, (ws.max_row or 0) + 1):
            wl = _to_float(ws.cell(r, 1).value)
            if wl is None:
                continue
            rows.append((
                sheet,
                wl,
                _to_float(ws.cell(r, 2).value),  # dBm
                _to_float(ws.cell(r, 3).value),  # nW
            ))
        repo.insert_measurement_spectra_points_bulk(measurement_id, rows)

    return measurement_id


def import_measurement_folder(repo, folder_path: Optional[str] = None, model: str = "") -> int:
    """
    ✅ 폴더 하위까지 재귀적으로 모든 *.xlsx Import
    - folder_path가 None/빈값이면 기본 RAW 루트(DEFAULT_RAW_ROOT)를 사용
    - stage/group_name은 raw_root 구조를 최대한 추론
    - model이 비어있으면 파일명/경로에서 자동 추출
    """
    if not folder_path:
        try:
            ps = normalize_and_fill_defaults(load_path_settings(repo))
            folder = Path((getattr(ps, "export_dir_measure", "") or "").strip() or ensure_default_raw_root())
        except Exception:
            folder = ensure_default_raw_root()
    else:
        folder = Path(folder_path)

    if not folder.exists():
        raise FileNotFoundError(str(folder))

    count = 0
    for f in sorted(folder.rglob("*.xlsx")):
        if f.name.startswith("~$"):
            continue

        # stage/model/group 추론:
        # 새 구조: raw_root\{stage}\{model}\{group}\file.xlsx
        # 기존 구조: raw_root\{stage}\{group}\file.xlsx
        stage = ""
        group_name = ""
        file_model = model or ""

        try:
            rel = f.relative_to(folder)
            parts = rel.parts
            if len(parts) >= 1:
                stage = parts[0]
            if len(parts) >= 2:
                # 두 번째 경로가 모델명 형태면 새 구조
                if re.fullmatch(r"F\d{3,4}\d{1,4}W", parts[1], flags=re.IGNORECASE):
                    file_model = file_model or parts[1]
                    if len(parts) >= 3:
                        group_name = parts[2]
                else:
                    # 기존 구조
                    group_name = parts[1]
            if len(parts) < 2:
                stage = folder.name
        except Exception:
            stage = f.parent.name

        import_measurement_xlsx(repo, str(f), stage=stage, group_name=group_name, model=file_model)
        count += 1

    return count
