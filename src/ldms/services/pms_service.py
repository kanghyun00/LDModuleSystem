# -*- coding: utf-8 -*-
# src/ldms/services/pms_service.py
"""
PMS service (logic only, no UI)

- Daily records saved as Excel (.xlsx) named by date: data/pms/YYYY-MM-DD.xlsx
- Weekly/Monthly reports exported to data/pms/
- Process name normalization: removes whitespace differences
  (e.g., "파이버 정렬" == "파이버정렬")

✅ 포함:
- Module No normalization: A1 == A001, e == E, etc.
- exe 안정화: 저장 경로는 src.ldms.config.get_pms_dir(), get_pms_daily_xlsx_path()만 사용
- Excel autosize + alignment 스타일 적용
- 삭제 기능(delete_daily_record_xlsx)
- ✅ 수정 기능(update_daily_record_xlsx): 삭제 + 재추가 방식
- 기간 로드(load_process_rows_between), 전체 레코드(fetch_daily_records_all)
- (참고) open_data_folder 포함
"""

from __future__ import annotations

import csv
import os
import re
import unicodedata
from pathlib import Path
from datetime import datetime, date, timedelta
from collections import defaultdict
from typing import List, Dict, Any, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment

from src.ldms.config import get_pms_daily_xlsx_path, get_pms_dir


KOR_HEADERS = ["날짜", "공정명", "개수", "모듈번호", "비고"]


# -------------------------
# ✅ Module normalization
# -------------------------
def normalize_module_no(raw: str) -> str:
    """
    모듈번호 정규화:
    - 대소문자 무시 (e == E)
    - 공백/구분기호 제거
    - (문자+숫자)면 숫자 최소 3자리 zero padding
      예) A1->A001, E1->E001, E110->E110
    - 숫자 4자리 이상이면 그대로 유지
    """
    if raw is None:
        return ""
    s = unicodedata.normalize("NFKC", str(raw)).strip()
    if not s:
        return ""
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"[^0-9A-Za-z]", "", s)
    s = s.upper()

    m = re.match(r"^([A-Z]+)(\d+)$", s)
    if not m:
        return s

    prefix, num_str = m.group(1), m.group(2)
    width = max(3, len(num_str))
    try:
        num_val = int(num_str)
    except Exception:
        return s
    return f"{prefix}{num_val:0{width}d}"


# -------------------------
# Path helpers (✅ config 기반)
# -------------------------
def ensure_pms_dir() -> Path:
    """
    ✅ PMS 폴더(data/pms) 생성 보장 (exe에서도 안정)
    """
    p = get_pms_dir()
    p.mkdir(parents=True, exist_ok=True)
    return p


def day_file_path_xlsx(d: date) -> Path:
    """
    날짜별 PMS 파일 경로(.xlsx)
    """
    return get_pms_daily_xlsx_path(d.strftime("%Y-%m-%d"))


def day_file_path_csv(d: date) -> Path:
    """
    (레거시) CSV 파일 경로
    """
    return ensure_pms_dir() / f"{d.strftime('%Y-%m-%d')}.csv"


# -------------------------
# Basic utils
# -------------------------
def today_str() -> str:
    return date.today().strftime("%Y-%m-%d")


def parse_date(s: str) -> Optional[date]:
    try:
        return datetime.strptime(str(s).strip(), "%Y-%m-%d").date()
    except Exception:
        return None


def safe_int(s: Any) -> Optional[int]:
    try:
        x = int(str(s).strip())
        if x < 0:
            return None
        return x
    except Exception:
        return None


def open_data_folder() -> None:
    """
    Windows: data/pms 폴더 열기
    """
    data_dir = ensure_pms_dir()
    try:
        os.startfile(str(data_dir))  # Windows only
    except Exception:
        pass


# -------------------------
# Excel style helpers
# -------------------------
def autosize_columns(ws, min_width: int = 10, max_width: int = 200) -> None:
    """
    셀 내용 기반 컬럼 폭 자동 조절 (줄바꿈 포함)
    """
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is None:
                continue
            s = str(v)
            s_len = max((len(line) for line in s.splitlines()), default=len(s))
            max_len = max(max_len, s_len)

        width = max(min_width, min(max_width, max_len + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def apply_excel_view_style(ws, center_all: bool = True, wrap_text: bool = True) -> None:
    """
    - 전체 셀 가로/세로 가운데정렬
    - wrap_text 적용
    - 컬럼 폭 자동 조절
    """
    if center_all:
        align = Alignment(horizontal="center", vertical="center", wrap_text=wrap_text)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = align

    autosize_columns(ws)


# -------------------------
# ✅ Process normalization (강화)
# -------------------------
_ALIAS_MAP: dict[str, str] = {
    # 필요하면 여기 추가:
    "어셈블리".upper(): "MIRROR1정렬",
    "MIRROR1": "MIRROR1정렬",
    "MIRROR1정렬": "MIRROR1정렬",
}


def normalize_process_name(name: str) -> str:
    """
    공정명 정규화(주간/월간/진행현황 공통 기준):
    - 유니코드 정규화(NFKC)
    - 공백 제거
    - 구분기호 제거
    - 영문 대문자 통일(PBS/VBG/LID 등)
    - alias 맵 적용
    """
    if name is None:
        return ""

    s = str(name)
    s = unicodedata.normalize("NFKC", s).strip()
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"[·•\-_\/]+", "", s)
    s = s.upper()
    return _ALIAS_MAP.get(s, s)


# -------------------------
# Daily read/write
# -------------------------
def write_day_records_xlsx(d: date, records: List[Dict[str, Any]]) -> Path:
    """
    records: [{"Process":.., "Qty":.., "Module":.., "Remark":..}, ...]
    """
    fp = day_file_path_xlsx(d)
    fp.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "일일공정"

    ws.append(KOR_HEADERS)
    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col in range(1, len(KOR_HEADERS) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.alignment = header_align

    for r in records:
        ws.append([
            d.strftime("%Y-%m-%d"),
            str(r.get("Process", "")).strip(),
            int(r.get("Qty", 0)) if safe_int(r.get("Qty", "0")) is not None else str(r.get("Qty", "")).strip(),
            str(r.get("Module", "")).strip(),
            str(r.get("Remark", "")).strip(),
        ])

    ws.freeze_panes = "A2"
    apply_excel_view_style(ws, center_all=True, wrap_text=True)
    wb.save(fp)
    return fp


def read_day_records_xlsx(d: date) -> List[Dict[str, str]]:
    fp = day_file_path_xlsx(d)
    if not fp.exists():
        return []

    wb = load_workbook(fp)
    ws = wb.active

    records: List[Dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        _date_val, proc, qty, module, remark = (row + (None, None, None, None, None))[:5]
        records.append({
            "Date": d.strftime("%Y-%m-%d"),
            "Process": (str(proc).strip() if proc is not None else ""),
            "Qty": (str(qty).strip() if qty is not None else ""),
            "Module": (str(module).strip() if module is not None else ""),
            "Remark": (str(remark).strip() if remark is not None else ""),
        })
    return records


def read_day_records_csv_legacy(d: date) -> List[Dict[str, str]]:
    fp = day_file_path_csv(d)
    if not fp.exists():
        return []
    rows: List[Dict[str, str]] = []
    with fp.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            rows.append({
                "Date": d.strftime("%Y-%m-%d"),
                "Process": str(r.get("Process", "")).strip(),
                "Qty": str(r.get("Qty", "")).strip(),
                "Module": str(r.get("Module", "")).strip(),
                "Remark": str(r.get("Remark", "")).strip(),
            })
    return rows


def load_records_for_day(d: date) -> List[Dict[str, str]]:
    if day_file_path_xlsx(d).exists():
        return read_day_records_xlsx(d)
    return read_day_records_csv_legacy(d)


# -------------------------
# Weekly/Monthly aggregation helpers
# -------------------------
def week_range_mon_fri(base_day: date) -> Tuple[date, date]:
    mon = base_day - timedelta(days=base_day.weekday())
    fri = mon + timedelta(days=4)
    return mon, fri


def month_range(base_day: date) -> Tuple[date, date]:
    start = base_day.replace(day=1)
    if start.month == 12:
        next_month = start.replace(year=start.year + 1, month=1, day=1)
    else:
        next_month = start.replace(month=start.month + 1, day=1)
    end = next_month - timedelta(days=1)
    return start, end


def iter_dates(start_d: date, end_d: date):
    cur = start_d
    while cur <= end_d:
        yield cur
        cur += timedelta(days=1)


def split_modules(raw: str) -> List[str]:
    """
    ✅ 엑셀 모듈번호 셀을 분해 + 정규화까지 수행
    """
    if raw is None:
        return []
    s = str(raw).strip()
    if not s:
        return []

    s = s.replace(";", ",").replace("\n", ",")
    parts = [x.strip() for x in s.split(",")]

    out: List[str] = []
    seen = set()

    for p in parts:
        if not p:
            continue
        m = normalize_module_no(p)
        if m and m not in seen:
            seen.add(m)
            out.append(m)

    return out


def build_detail_rows(start_d: date, end_d: date) -> List[Dict[str, Any]]:
    detail: List[Dict[str, Any]] = []
    for d in iter_dates(start_d, end_d):
        recs = load_records_for_day(d)
        for r in recs:
            qty = safe_int(r.get("Qty", "0")) or 0

            proc_raw = r.get("Process", "")
            proc_norm = normalize_process_name(proc_raw)

            detail.append({
                "Date": d.strftime("%Y-%m-%d"),
                "Process": proc_norm,
                "ProcessRaw": proc_raw,
                "Qty": qty,
                "Module": str(r.get("Module", "")).strip(),
                "Remark": str(r.get("Remark", "")).strip(),
            })
    return detail


def build_summary_unique_modules(detail_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    qty_sum = defaultdict(int)
    days_set = defaultdict(set)
    modules_set = defaultdict(set)

    for r in detail_rows:
        p = r.get("Process") or "(미지정)"
        qty_sum[p] += int(r.get("Qty", 0))
        days_set[p].add(r.get("Date"))

        for m in split_modules(r.get("Module", "")):
            modules_set[p].add(m)

    summary: List[Dict[str, Any]] = []
    for p in sorted(qty_sum.keys()):
        modules_sorted = sorted(modules_set[p])
        summary.append({
            "Process": p,
            "UniqueModules": len(modules_sorted),
            "QtySum": qty_sum[p],
            "DaysCount": len(days_set[p]),
            "Modules": ", ".join(modules_sorted),
        })
    return summary


# -------------------------
# Export weekly/monthly
# -------------------------
def export_weekly_xlsx(
    start_d: date,
    end_d: date,
    detail_rows: List[Dict[str, Any]],
    summary_rows: List[Dict[str, Any]],
) -> Path:
    data_dir = ensure_pms_dir()
    fp = data_dir / f"주간리포트_{start_d.strftime('%Y-%m-%d')}_to_{end_d.strftime('%Y-%m-%d')}.xlsx"

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "주간상세"
    ws1.append(KOR_HEADERS)

    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, 6):
        c = ws1.cell(row=1, column=col)
        c.font = header_font
        c.alignment = header_align
    ws1.freeze_panes = "A2"

    for r in detail_rows:
        ws1.append([r["Date"], r["Process"], r["Qty"], r["Module"], r["Remark"]])

    apply_excel_view_style(ws1, center_all=True, wrap_text=True)

    ws2 = wb.create_sheet("주간요약")
    ws2.append(["공정명", "완료수(고유모듈)", "Qty합(참고)", "진행일수(월~금)", "모듈(고유목록)"])
    for col in range(1, 6):
        c = ws2.cell(row=1, column=col)
        c.font = header_font
        c.alignment = header_align
    ws2.freeze_panes = "A2"

    for r in summary_rows:
        ws2.append([r["Process"], r["UniqueModules"], r["QtySum"], r["DaysCount"], r["Modules"]])

    apply_excel_view_style(ws2, center_all=True, wrap_text=True)

    wb.save(fp)
    return fp


def export_monthly_xlsx(
    start_d: date,
    end_d: date,
    detail_rows: List[Dict[str, Any]],
    summary_rows: List[Dict[str, Any]],
) -> Path:
    data_dir = ensure_pms_dir()
    fp = data_dir / f"월간리포트_{start_d.strftime('%Y-%m')}.xlsx"

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "월간상세"
    ws1.append(KOR_HEADERS)

    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, 6):
        c = ws1.cell(row=1, column=col)
        c.font = header_font
        c.alignment = header_align
    ws1.freeze_panes = "A2"

    for r in detail_rows:
        ws1.append([r["Date"], r["Process"], r["Qty"], r["Module"], r["Remark"]])

    apply_excel_view_style(ws1, center_all=True, wrap_text=True)

    ws2 = wb.create_sheet("월간요약")
    ws2.append(["공정명", "완료수(고유모듈)", "Qty합(참고)", "진행일수", "모듈(고유목록)"])
    for col in range(1, 6):
        c = ws2.cell(row=1, column=col)
        c.font = header_font
        c.alignment = header_align
    ws2.freeze_panes = "A2"

    for r in summary_rows:
        ws2.append([r["Process"], r["UniqueModules"], r["QtySum"], r["DaysCount"], r["Modules"]])

    apply_excel_view_style(ws2, center_all=True, wrap_text=True)

    wb.save(fp)
    return fp


# -------------------------
# Append / Delete / ✅ Update
# -------------------------
def append_daily_record_xlsx(work_date_str: str, process: str, qty: int, module: str, remark: str) -> Path:
    """
    특정 날짜의 PMS 일일 엑셀(data/pms/YYYY-MM-DD.xlsx)에 한 줄을 '누적' 저장한다.
    ✅ module은 저장 전에 정규화하여 A1/A001 혼재를 줄임
    """
    d = parse_date(work_date_str)
    if d is None:
        raise ValueError(f"날짜 형식 오류(YYYY-MM-DD): {work_date_str}")

    records = load_records_for_day(d)

    proc_norm = normalize_process_name(process)

    modules_norm = split_modules(module)
    module_csv_norm = ", ".join(modules_norm)

    records.append({
        "Process": proc_norm,
        "Qty": str(int(qty)),
        "Module": module_csv_norm,
        "Remark": str(remark).strip(),
    })

    return write_day_records_xlsx(d, records)


def _norm(v):
    return "" if v is None else str(v).strip()


def _norm_modules(s):
    modules = split_modules("" if s is None else str(s))
    return ",".join(modules)


def delete_daily_record_xlsx(
    work_date: str,
    process: str,
    qty: int,
    module_csv: str,
    remark: str,
) -> bool:
    """
    날짜별 PMS 파일(data/pms/YYYY-MM-DD.xlsx)에서
    (공정명, 개수, 모듈번호, 비고)가 일치하는 첫 행 삭제
    """
    try:
        xlsx_path = get_pms_daily_xlsx_path(work_date)
        if not xlsx_path.exists():
            return False

        wb = load_workbook(xlsx_path)
        ws = wb.active

        target = (
            _norm(process),
            _norm(qty),
            _norm_modules(module_csv),
            _norm(remark),
        )

        for r in range(2, ws.max_row + 1):
            row = (
                _norm(ws.cell(r, 2).value),
                _norm(ws.cell(r, 3).value),
                _norm_modules(ws.cell(r, 4).value),
                _norm(ws.cell(r, 5).value),
            )
            if row == target:
                ws.delete_rows(r, 1)
                apply_excel_view_style(ws, center_all=True, wrap_text=True)
                wb.save(xlsx_path)
                return True

        return False
    except Exception:
        return False


def update_daily_record_xlsx(
    work_date: str,
    old_process: str,
    old_qty: int,
    old_module_csv: str,
    old_remark: str,
    new_process: str,
    new_qty: int,
    new_module_csv: str,
    new_remark: str,
) -> bool:
    """
    PMS 엑셀에서 기존 행을 찾아 삭제한 뒤, 수정된 값으로 다시 append
    - 반환값: 기존 행을 찾았으면 True, 못 찾았으면 False
    """
    ok = delete_daily_record_xlsx(work_date, old_process, old_qty, old_module_csv, old_remark)
    append_daily_record_xlsx(work_date, new_process, new_qty, new_module_csv, new_remark)
    return ok


# -------------------------
# 기간 로딩 / 집계
# -------------------------
def _parse_date(s: str) -> date:
    return datetime.strptime(str(s).strip(), "%Y-%m-%d").date()


def _week_range_monday(base: date):
    """월요일~일요일"""
    monday = base - timedelta(days=base.weekday())
    sunday = monday + timedelta(days=6)
    return monday, sunday


def _iter_pms_daily_files(start: date, end: date):
    """data/pms/YYYY-MM-DD.xlsx 파일을 날짜 범위로 순회"""
    pms_dir = get_pms_dir()
    d = start
    while d <= end:
        path = pms_dir / f"{d.strftime('%Y-%m-%d')}.xlsx"
        if path.exists():
            yield d, path
        d += timedelta(days=1)


def _header_index_map(ws) -> Dict[str, int]:
    """
    PMS Excel 헤더를 읽어 컬럼 위치를 반환합니다.

    지원 형식:
    구형: 날짜 | 공정명 | 개수 | 모듈번호 | 비고
    신형: 날짜 | 모델명 | 공정명 | 개수 | 모듈번호 | 비고
    """
    aliases = {
        "date": {"날짜", "date", "work_date"},
        "model": {"모델", "모델명", "model", "model_code", "product_model"},
        "process": {"공정명", "공정", "process", "process_name"},
        "qty": {"개수", "수량", "qty", "quantity"},
        "module": {"모듈번호", "모듈", "module", "module_no", "modules"},
        "remark": {"비고", "remark", "note"},
    }

    result: Dict[str, int] = {}

    for col_idx, cell in enumerate(ws[1], start=1):
        value = "" if cell.value is None else str(cell.value).strip()
        if not value:
            continue

        normalized = value.replace(" ", "").lower()

        for key, candidates in aliases.items():
            if normalized in {
                candidate.replace(" ", "").lower()
                for candidate in candidates
            }:
                result[key] = col_idx
                break

    # 헤더를 인식하지 못한 레거시 파일은 고정 5열로 처리합니다.
    if "process" not in result:
        result["process"] = 2
    if "qty" not in result:
        result["qty"] = 3
    if "module" not in result:
        result["module"] = 4
    if "remark" not in result:
        result["remark"] = 5

    return result


def _cell_text(ws, row_idx: int, col_idx: Optional[int]) -> str:
    if not col_idx or col_idx < 1:
        return ""

    value = ws.cell(row=row_idx, column=col_idx).value
    return "" if value is None else str(value).strip()


def _cell_int(ws, row_idx: int, col_idx: Optional[int]) -> int:
    value = _cell_text(ws, row_idx, col_idx)

    if not value:
        return 0

    try:
        return max(0, int(float(value)))
    except Exception:
        return 0


def load_process_rows_between(start_date: str, end_date: str) -> list[dict]:
    """
    PMS Excel 기간 자료를 읽습니다.

    구형 5열과 신형 6열을 모두 지원합니다.
    모델 정보가 없는 구형 행은 model을 빈 문자열로 반환하며,
    실제 모델 판정은 ProcessTab에서 DB 모듈 매핑으로 수행합니다.
    """
    start = _parse_date(start_date)
    end = _parse_date(end_date)
    rows: list[dict] = []

    for d, path in _iter_pms_daily_files(start, end):
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            header_map = _header_index_map(ws)

            date_col = header_map.get("date")
            model_col = header_map.get("model")
            process_col = header_map.get("process")
            qty_col = header_map.get("qty")
            module_col = header_map.get("module")
            remark_col = header_map.get("remark")

            for row_idx in range(2, ws.max_row + 1):
                process = _cell_text(ws, row_idx, process_col)
                modules_csv = _cell_text(ws, row_idx, module_col)

                if not process and not modules_csv:
                    continue

                work_date = _cell_text(ws, row_idx, date_col) or d.strftime("%Y-%m-%d")
                raw_model = _cell_text(ws, row_idx, model_col)
                qty = _cell_int(ws, row_idx, qty_col)
                remark = _cell_text(ws, row_idx, remark_col)

                proc_norm = normalize_process_name(process)
                modules_csv_norm = ",".join(split_modules(modules_csv))

                rows.append({
                    "work_date": work_date,
                    "process": process,
                    "process_norm": proc_norm,
                    "qty": qty,
                    "modules_csv": modules_csv_norm,
                    "remark": remark,
                    "model": raw_model,
                    "model_code": raw_model,
                    "product_model": raw_model,
                })

            wb.close()

        except Exception:
            # 특정 파일이 손상되었거나 열려 있어도 다른 파일은 계속 읽습니다.
            continue

    return rows


def summarize_week(base_date_str: str) -> dict:
    """
    기준일이 속한 주(월~일) 집계 결과를 반환.
    """
    base = _parse_date(base_date_str)
    start, end = _week_range_monday(base)
    rows = load_process_rows_between(start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d"))

    process_qty_sum: dict[str, int] = {}
    process_unique_set: dict[str, set[str]] = {}

    for row in rows:
        proc = row.get("process_norm") or normalize_process_name(row.get("process", ""))
        qty = int(row.get("qty") or 0)

        process_qty_sum[proc] = process_qty_sum.get(proc, 0) + qty

        if proc not in process_unique_set:
            process_unique_set[proc] = set()

        for m in split_modules(row.get("modules_csv", "")):
            process_unique_set[proc].add(m)

    process_unique_modules = {k: len(v) for k, v in process_unique_set.items()}

    return {
        "week_start": start.strftime("%Y-%m-%d"),
        "week_end": end.strftime("%Y-%m-%d"),
        "process_qty_sum": process_qty_sum,
        "process_unique_modules": process_unique_modules,
        "total_rows": len(rows),
    }


# ============================================================
# ✅ 진행현황(모듈별/공정별 대기)용: 전체 일일공정 레코드 추출
# ============================================================
def fetch_daily_records_all(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
) -> List[Dict[str, str]]:
    """
    PMS data/pms/YYYY-MM-DD.xlsx 전부(또는 기간) 읽어서
    진행현황 계산용 레코드로 반환

    return:
      [
        {"date":"2026-02-04", "process":"어셈블리", "module_no":"A001"},
        ...
      ]

    - 모듈번호는 split_modules에서 정규화까지 완료
    """
    pms_dir = get_pms_dir()

    files = []
    for fp in pms_dir.glob("*.xlsx"):
        stem = fp.stem  # YYYY-MM-DD
        d = parse_date(stem)
        if d is None:
            continue
        files.append((d, fp))

    files.sort(key=lambda x: x[0])

    sd = parse_date(start_date) if start_date else None
    ed = parse_date(end_date) if end_date else None

    out: List[Dict[str, str]] = []

    for d, fp in files:
        if sd and d < sd:
            continue
        if ed and d > ed:
            continue

        try:
            wb = load_workbook(fp)
            ws = wb.active
        except Exception:
            continue

        for r in range(2, ws.max_row + 1):
            proc = ws.cell(r, 2).value
            modules_csv = ws.cell(r, 4).value

            proc_raw = "" if proc is None else str(proc).strip()
            if not proc_raw:
                continue

            modules = split_modules("" if modules_csv is None else str(modules_csv))
            if not modules:
                continue

            for m in modules:
                out.append({
                    "date": d.strftime("%Y-%m-%d"),
                    "process": proc_raw,
                    "module_no": m,   # ✅ 이미 정규화됨
                })

    return out
