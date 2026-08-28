from __future__ import annotations

from pathlib import Path
from typing import Dict

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font

from src.ldms.config import get_reports_dir


def export_modules_report(
    start_date: str,
    end_date: str,
    proc_to_modules: Dict[str, set],
    kind: str = "주간",
) -> Path:
    """
    공정별 유니크 모듈 집계를 엑셀로 저장
    컬럼: 공정명 / 총 진행 모듈 수(중복제외) / 진행된 모듈 번호

    ✅ 저장 위치: config.get_reports_dir() (exe에서도 Documents or portable data)
    """
    safe_kind = str(kind).strip() or "리포트"
    filename = f"{safe_kind}리포트_{start_date}_to_{end_date}.xlsx"
    out_path = get_reports_dir() / filename

    wb = Workbook()
    ws = wb.active
    ws.title = f"{safe_kind}리포트"

    headers = ["공정명", "총 진행 모듈 수(중복제외)", "진행된 모듈 번호"]
    ws.append(headers)

    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, 4):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.alignment = header_align

    items = sorted(proc_to_modules.items(), key=lambda x: len(x[1]), reverse=True)
    for proc, modset in items:
        mods = sorted(list(modset))
        ws.append([proc, len(modset), ", ".join(mods)])

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.alignment = center

    for col_idx in range(1, 4):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is None:
                continue
            s = str(v)
            s_len = max((len(line) for line in s.splitlines()), default=len(s))
            max_len = max(max_len, s_len)
        width = max(10, min(200, max_len + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{max(1, ws.max_row)}"

    wb.save(out_path)
    return out_path
